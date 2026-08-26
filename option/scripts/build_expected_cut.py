# -*- coding: utf-8 -*-
"""T-031 L2+L3：企业绩效等级与 expected_cut（优先用地方实际减排清单）

证据优先级（高 -> 低）：
  1. 邢台市《重污染天气工业源减排清单》：逐企业、逐预警级别的**绝对投料量基准与目标**，
     可算出点值 expected_cut。口径为"投料量和拉延量（车速×厚度×板宽）"，与熔窑负荷同物理量。
  2. 邢台市逐季度 D升C / 降级公示：给出等级变动（原等级→新等级）及日期。
  3. 河北省生态环境厅年度 A/B/引领性清单（单一 vintage 2024-04-28）。
  4. 技术指南 §19（五）比例表：仅在上述均无时作为兜底，且 unlisted 只能出 [C,D] 区间。

比例表来源：生态环境部《重污染天气重点行业应急减排措施制定技术指南（2020 修订版）》。
**该文件处于换版期（环办大气〔2026〕2 号新增 A+ 级），每次使用前须复核是否已有新版。**

输入：data/public_official/{reduction_list,perf_grade}/、glass_line_port_mapping.csv、alert_calendar_邢台市.csv
输出：data/public_official/expected_cut.csv
"""
import csv, glob, io, os, re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PO = os.path.join(ROOT, 'data/public_official')
PG = os.path.join(PO, 'perf_grade')
RL = os.path.join(PO, 'reduction_list')

GUIDE = {'A': (set(), 0.0), 'leading': (set(), 0.0),
         'B': ({'orange', 'red'}, 0.10),
         'C': ({'orange', 'red'}, 0.20),
         'D': ({'yellow', 'orange', 'red'}, 0.30)}
LEVEL_COL = {'red': 5, 'orange': 6, 'yellow': 7}     # 减排清单中三级预警措施所在列
norm = lambda s: re.sub(r'[\s()（）]', '', s or '')
# "原料日投料量由960吨降低到768吨及以下"
FEED = re.compile(r'投料量?\s*由\s*([\d.]+)\s*吨\s*(?:降低|降)\s*到\s*([\d.]+)\s*吨')


def read_rows(fn, cap=40000):
    try:
        if fn.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(fn, read_only=True, data_only=True)
            out = []
            for sn in wb.sheetnames:
                for i, r in enumerate(wb[sn].iter_rows(values_only=True)):
                    if i > cap: break
                    out.append([('' if c is None else str(c)).strip() for c in r])
            return out
        import xlrd
        wb = xlrd.open_workbook(fn)
        return [[str(c.value).strip() for c in s.row(i)]
                for s in wb.sheets() for i in range(min(s.nrows, cap))]
    except Exception as e:
        print('  [warn] 读取失败 %s: %s' % (os.path.basename(fn), e))
        return []


def local_list():
    """企业 -> dict(grade, long_idle, cuts={level: (basis, target, cut)})"""
    fn = os.path.join(RL, 'xingtai_industrial_reduction_2024.xlsx')
    if not os.path.exists(fn):
        return {}
    out = {}
    for r in read_rows(fn)[1:]:
        if len(r) < 8 or not r[2]:
            continue
        k = norm(r[2])
        d = out.setdefault(k, dict(name=r[2], county=r[3], grades=set(), long_idle=False, cuts={}))
        g = (r[4] or '').strip()
        if '长期停产' in g:
            d['long_idle'] = True
        gg = re.sub(r'^(地方|民生豁免-)', '', g)
        if gg in ('A', 'B', 'C', 'D'):
            d['grades'].add(gg)
        for lv, col in LEVEL_COL.items():
            m = FEED.search(r[col] or '')
            if m:
                a, b = float(m.group(1)), float(m.group(2))
                if a > 0:
                    d['cuts'][lv] = (a, b, round(1 - b / a, 4))
    return out


def transitions():
    out = []
    for fn in sorted(glob.glob(os.path.join(PG, '*.xls*'))):
        m = re.match(r'(\d{4}-\d{2}-\d{2})_', os.path.basename(fn))
        if not m:
            continue
        for r in read_rows(fn):
            cells = [c for c in r if c]
            if not any('玻璃' in c for c in cells):
                continue
            nm = next((c for c in cells if c.endswith(('有限公司', '股份有限公司', '玻璃厂'))), None)
            gs = [c for c in cells if c in ('A', 'B', 'C', 'D')]
            if nm and len(gs) >= 2:
                out.append(dict(date=m.group(1), name=nm, frm=gs[0], to=gs[1]))
    return out


def prov_grades():
    g = {}
    for tag, fn in (('A', 'prov_A_2024.xlsx'), ('B', 'prov_B_2024.xlsx'), ('leading', 'prov_leading_2024.xlsx')):
        for r in read_rows(os.path.join(PG, fn))[1:]:
            if len(r) > 1 and r[1]:
                g[norm(r[1])] = tag
    return g


def windows(city='邢台市'):
    rows = list(csv.DictReader(io.open(os.path.join(PO, 'alert_calendar_%s.csv' % city), encoding='utf-8-sig')))
    rows.sort(key=lambda r: (r['pubdate'], int(r['act_id'] or 0)))
    win, cur = [], None
    for r in rows:
        if r['action'] in ('raise', 'upgrade'):
            if cur is None:
                cur = dict(start=r['pubdate'], level=r['level_en'] or 'orange')
            elif r['level_en']:
                cur['level'] = r['level_en']
        elif r['action'] == 'lift' and cur:
            cur['end'] = r['pubdate']; win.append(cur); cur = None
    return win


def resolve(ent, day, loc, prov, trans):
    """返回 (grade, source, long_idle)。"""
    k = norm(ent)
    hist = sorted([t for t in trans if norm(t['name']) == k and t['date'] <= day], key=lambda x: x['date'])
    if k in loc and loc[k]['grades']:
        g = sorted(loc[k]['grades'])[0]
        return g, '邢台市工业源减排清单', loc[k]['long_idle']
    if hist:
        return hist[-1]['to'], '%s 等级变动 %s→%s' % (hist[-1]['date'], hist[-1]['frm'], hist[-1]['to']), False
    if k in prov:
        return prov[k], '省厅 2024 年度清单', False
    return None, '未列入任何清单（只能为 C 或 D）', False


def main():
    loc, prov, trans = local_list(), prov_grades(), transitions()
    print('本地减排清单企业 %d 家；省厅清单 %d 家；玻璃等级变动 %d 条' % (len(loc), len(prov), len(trans)))
    mp = [m for m in csv.DictReader(io.open(os.path.join(PO, 'glass_line_port_mapping.csv'), encoding='utf-8-sig'))
          if int(m['float_lines']) > 0]
    wins = windows()
    out = []
    for w in wins:
        for m in mp:
            k = norm(m['enterprise'])
            g, src, idle = resolve(m['enterprise'], w['start'], loc, prov, trans)
            basis = target = ''
            if k in loc and w['level'] in loc[k]['cuts']:
                b, t, c = loc[k]['cuts'][w['level']]
                lo = hi = c; basis, target, method = b, t, '清单绝对吨数'
            elif g:
                trig, cut = GUIDE[g]
                lo = hi = (cut if w['level'] in trig else 0.0); method = '技术指南比例'
            else:
                vals = [(GUIDE[c][1] if w['level'] in GUIDE[c][0] else 0.0) for c in ('C', 'D')]
                lo, hi, method = min(vals), max(vals), '技术指南区间(C|D)'
            out.append(dict(window_start=w['start'], window_end=w['end'], level=w['level'],
                            enterprise=m['enterprise'], float_lines=m['float_lines'],
                            float_tpd=int(m['float_tpd']), resolvable_lines=m.get('resolvable_lines', ''),
                            grade=g or 'C|D', grade_source=src, long_idle_flag=int(idle),
                            feed_basis_tpd=basis, feed_target_tpd=target,
                            expected_cut_lo=round(lo, 4), expected_cut_hi=round(hi, 4),
                            method=method, is_interval=int(lo != hi)))
    with io.open(os.path.join(PO, 'expected_cut.csv'), 'w', encoding='utf-8-sig', newline='') as f:
        wr = csv.DictWriter(f, fieldnames=list(out[0].keys())); wr.writeheader(); wr.writerows(out)

    tot = sum(int(m['float_tpd']) for m in mp)
    latest = wins[-1]['start']
    print('\n截至最近窗口 %s：' % latest)
    print('%-26s %5s %7s %-5s %-22s %s' % ('企业', '浮法线', 't/d', '等级', '依据', '长期停产'))
    unres = 0
    for m in mp:
        g, src, idle = resolve(m['enterprise'], latest, loc, prov, trans)
        if not g: unres += int(m['float_tpd'])
        print('%-26s %5s %7s %-5s %-22s %s' % (m['enterprise'][:26], m['float_lines'], m['float_tpd'],
                                               g or 'C|D', src[:22], '是' if idle else ''))
    pt = sum(1 for r in out if not r['is_interval'])
    ab = sum(1 for r in out if r['method'] == '清单绝对吨数')
    print('\n输出 %d 行；点值 %d（%.1f%%），其中来自清单绝对吨数 %d（%.1f%%）'
          % (len(out), pt, 100.0*pt/len(out), ab, 100.0*ab/len(out)))
    print('等级仍不可分辨的浮法产能：%d / %d t/d = %.1f%%（前值 76.6%%）' % (unres, tot, 100.0*unres/tot))


if __name__ == '__main__':
    main()
